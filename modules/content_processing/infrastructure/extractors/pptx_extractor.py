from pathlib import Path
from pptx import Presentation


class PptxContentExtractor:
    def extract_markdown(self, path: Path) -> tuple[str, int]:
        """Extract clean Markdown and page count (number of slides) from a PPTX file."""
        prs = Presentation(str(path))
        slides_text: list[str] = []

        for idx, slide in enumerate(prs.slides, start=1):
            slide_lines: list[str] = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for paragraph in shape.text_frame.paragraphs:
                        text = paragraph.text.strip()
                        if text:
                            slide_lines.append(text)
            if slide_lines:
                slides_text.append(f"## Diapositiva {idx}\n\n" + "\n\n".join(slide_lines))

        markdown = "\n\n---\n\n".join(slides_text).strip()
        page_count = max(1, len(prs.slides))
        return markdown, page_count
