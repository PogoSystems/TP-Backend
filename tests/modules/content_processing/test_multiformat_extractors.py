import io
from pathlib import Path
import zipfile
import pytest
from pptx import Presentation

from modules.content_processing.application.services.document_service import ALLOWED_CONTENT_TYPES
from modules.content_processing.application.services.parsing.content_preparation_service import (
    ContentPreparationService,
)
from modules.content_processing.domain.aggregates.content_document import ContentDocumentAggregate
from modules.content_processing.domain.value_objects.prepared_content import PreparedDocument
from modules.content_processing.infrastructure.extractors.document_extractor import (
    DocumentContentExtractor,
)
from modules.content_processing.infrastructure.extractors.docx_extractor import DocxContentExtractor
from modules.content_processing.infrastructure.extractors.pptx_extractor import PptxContentExtractor
from modules.content_processing.infrastructure.storage.local_document_store import LocalDocumentStore
from shared.exceptions import InvalidFileTypeError


def _create_sample_docx(path: Path) -> None:
    """Helper to write a minimal valid docx with a heading and paragraph."""
    content_types = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
        '<Default Extension="xml" ContentType="application/xml"/>'
        '<Override PartName="/word/document.xml" '
        'ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>'
        '</Types>'
    )
    document_xml = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        '<w:body>'
        '<w:p><w:pPr><w:pStyle w:val="Heading1"/></w:pPr><w:r><w:t>Introduccion al Curso</w:t></w:r></w:p>'
        '<w:p><w:r><w:t>Contenido detallado en formato Word.</w:t></w:r></w:p>'
        '</w:body>'
        '</w:document>'
    )
    with zipfile.ZipFile(path, "w") as zf:
        zf.writestr("[Content_Types].xml", content_types)
        zf.writestr("word/document.xml", document_xml)


def _create_sample_pptx(path: Path) -> None:
    """Helper to write a minimal valid pptx with 2 slides."""
    prs = Presentation()
    slide1 = prs.slides.add_slide(prs.slide_layouts[0])
    slide1.shapes.title.text = "Tema 1: Arquitectura"

    slide2 = prs.slides.add_slide(prs.slide_layouts[1])
    slide2.shapes.title.text = "Tema 2: Backend"
    for shape in slide2.shapes:
        if shape.has_text_frame and shape != slide2.shapes.title:
            shape.text_frame.text = "Puntos clave de la presentacion."
    prs.save(str(path))


def test_docx_extractor_extracts_markdown(tmp_path: Path) -> None:
    docx_file = tmp_path / "test.docx"
    _create_sample_docx(docx_file)

    extractor = DocxContentExtractor()
    markdown, page_count = extractor.extract_markdown(docx_file)

    assert "Introduccion al Curso" in markdown
    assert "Contenido detallado" in markdown
    assert page_count >= 1


def test_pptx_extractor_extracts_markdown(tmp_path: Path) -> None:
    pptx_file = tmp_path / "test.pptx"
    _create_sample_pptx(pptx_file)

    extractor = PptxContentExtractor()
    markdown, page_count = extractor.extract_markdown(pptx_file)

    assert "## Diapositiva 1" in markdown
    assert "Tema 1: Arquitectura" in markdown
    assert "## Diapositiva 2" in markdown
    assert page_count == 2


def test_document_extractor_black_box_routing(tmp_path: Path) -> None:
    extractor = DocumentContentExtractor()

    docx_file = tmp_path / "notes.docx"
    _create_sample_docx(docx_file)
    md_docx, pages_docx = extractor.extract_markdown(docx_file)
    assert "Introduccion al Curso" in md_docx
    assert pages_docx >= 1

    pptx_file = tmp_path / "slides.pptx"
    _create_sample_pptx(pptx_file)
    md_pptx, pages_pptx = extractor.extract_markdown(pptx_file)
    assert "## Diapositiva 1" in md_pptx
    assert pages_pptx == 2

    # Unsupported format raises InvalidFileTypeError
    unsupported = tmp_path / "audio.mp3"
    unsupported.write_bytes(b"fake audio")
    with pytest.raises(InvalidFileTypeError):
        extractor.extract_markdown(unsupported)


def test_local_document_store_preserves_extensions(tmp_path: Path) -> None:
    store = LocalDocumentStore(base_path=tmp_path)

    docx_file = tmp_path / "source.docx"
    docx_file.write_bytes(b"fake docx")
    docx_key = store.store_document(docx_file)
    assert docx_key.endswith(".docx")
    assert (tmp_path / docx_key).exists()

    pptx_file = tmp_path / "source.pptx"
    pptx_file.write_bytes(b"fake pptx")
    pptx_key = store.store_document(pptx_file)
    assert pptx_key.endswith(".pptx")
    assert (tmp_path / pptx_key).exists()

    # store_pdf backward compatibility
    pdf_file = tmp_path / "source.pdf"
    pdf_file.write_bytes(b"fake pdf")
    pdf_key = store.store_pdf(pdf_file)
    assert pdf_key.endswith(".pdf")
    assert (tmp_path / pdf_key).exists()


def test_preparation_service_handles_multiformat(tmp_path: Path) -> None:
    store = LocalDocumentStore(base_path=tmp_path)
    service = ContentPreparationService(document_store=store)

    docx_file = tmp_path / "tema1.docx"
    _create_sample_docx(docx_file)
    prep_docx = service.prepare_document(source_path=docx_file, title="Tema 1")
    assert isinstance(prep_docx, PreparedDocument)
    assert prep_docx.raw.document_type == "docx"
    assert len(prep_docx.sections) > 0

    pptx_file = tmp_path / "slides.pptx"
    _create_sample_pptx(pptx_file)
    prep_pptx = service.prepare_document(source_path=pptx_file, title="Slides")
    assert isinstance(prep_pptx, PreparedDocument)
    assert prep_pptx.raw.document_type == "pptx"
    assert len(prep_pptx.sections) > 0


def test_allowed_content_types_contains_pptx_and_docx() -> None:
    assert "application/vnd.openxmlformats-officedocument.wordprocessingml.document" in ALLOWED_CONTENT_TYPES
    assert "application/vnd.openxmlformats-officedocument.presentationml.presentation" in ALLOWED_CONTENT_TYPES
    assert "application/pdf" in ALLOWED_CONTENT_TYPES
    assert "pptx" in ContentDocumentAggregate.ALLOWED_TYPES
    assert "docx" in ContentDocumentAggregate.ALLOWED_TYPES
    assert "pdf" in ContentDocumentAggregate.ALLOWED_TYPES
